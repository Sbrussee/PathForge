"""End-to-end smoke coverage for precomputed SISH slide retrieval.

The production SISH path can consume patch indices produced during a prior
precomputation step.  This smoke test uses that supported input form so it
exercises task orchestration, index persistence, and ranking without requiring
the optional VQ-VAE and codebook assets.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

from pathforge.core.datasets.bag_dataset import SlideRetrievalBagDataset
from pathforge.core.experiments.combinations import ComboConfig
from pathforge.core.experiments.combo_ids import build_tiling_id
from pathforge.core.tasks.slide_retrieval import SlideRetrievalTask
from pathforge.slide_retrieval.representation_strategies.types import (
    RetrievalRepresentation,
)
from pathforge.slide_retrieval.types import RetrievalItemMetadata


class _Sample:
    """Minimal slide-retrieval sample used when representations are pre-cached."""

    def __init__(self, sample_id: str) -> None:
        self.sample_id = sample_id


class _PrecomputedSISHBagDataset(SlideRetrievalBagDataset):
    """Small in-memory dataset sufficient for the task's cache-first path."""

    def __init__(self, *, name: str, tiling_id: str, sample_ids: list[str]) -> None:
        self._name = name
        self.tiling_id = tiling_id
        self.aggregation_level = "slide"
        self._samples = [_Sample(sample_id) for sample_id in sample_ids]

    @property
    def num_bags(self) -> int:
        return len(self._samples)

    def get_sample(self, index: int) -> _Sample:
        return self._samples[index]

    def get_feature_level(self) -> str:
        return "patch"

    def get_feature_level_reason(self) -> str:
        return "precomputed SISH patch representations"


def _combo() -> ComboConfig:
    """Return a patch-vector/SISH-compatible task combination."""
    return ComboConfig(
        tile_px=224,
        tile_px_params={},
        tile_mpp=1.0,
        tile_mpp_params={},
        feature_extraction="resnet18",
        feature_extraction_params={},
        retrieval_representation="hshr-features",
        retrieval_representation_params={},
        search_strategy="sish",
        search_strategy_params={
            "k": 2,
            "seed_fanout_t": 1,
            "pre_step": 1,
            "succ_step": 1,
            "hamming_thr": 8,
        },
    )


def _representation(*, sample_id: str, patient_id: str, patch_index: int) -> RetrievalRepresentation:
    """Build one valid precomputed SISH representation."""
    return RetrievalRepresentation(
        sample_id=sample_id,
        data=np.asarray([[1.0, 1.0, 1.0, 1.0]], dtype=np.float32),
        metadata=RetrievalItemMetadata(category="tumor", patient_id=patient_id),
        additional_data={
            "sish_patch_indices": np.asarray([patch_index], dtype=np.int64),
            "selected_coords": np.asarray([[0, 0]], dtype=np.int32),
        },
    )


def _register_retrieval_strategies() -> None:
    """Register production retrieval strategies for direct task execution."""
    from pathforge.slide_retrieval.representation_strategies.registry import (
        import_representation_strategy_modules,
    )
    from pathforge.slide_retrieval.search_strategies.registry import (
        import_search_strategy_modules,
    )

    import_representation_strategy_modules()
    import_search_strategy_modules()


@pytest.mark.smoke
def test_smoke_slide_retrieval_sish_uses_precomputed_indices(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """SISH indexes precomputed patches and ranks the matching reference first."""
    _register_retrieval_strategies()
    combo = _combo()
    tiling_id = build_tiling_id(combo)

    reference_dataset = _PrecomputedSISHBagDataset(
        name="reference",
        tiling_id=tiling_id,
        sample_ids=["reference-match", "reference-other"],
    )
    query_dataset = _PrecomputedSISHBagDataset(
        name="query", tiling_id=tiling_id, sample_ids=["query-slide"]
    )
    representations = {
        "reference": [
            _representation(
                sample_id="reference-match", patient_id="patient-a", patch_index=19
            ),
            _representation(
                sample_id="reference-other", patient_id="patient-b", patch_index=21
            ),
        ],
        "query": [
            _representation(
                sample_id="query-slide", patient_id="patient-query", patch_index=20
            )
        ],
    }

    def _precomputed_cache(
        self,
        *,
        bag_dataset,
        representation_id,
        aggregation_level,
        exclusion_level,
    ):
        del self, representation_id, aggregation_level, exclusion_level
        return representations[bag_dataset._name], None

    monkeypatch.setattr(
        SlideRetrievalTask, "_collect_existing_representations", _precomputed_cache
    )

    sish_dir = tmp_path / "sish"
    cfg = SimpleNamespace(
        experiment=SimpleNamespace(
            aggregation_level="slide",
            num_workers=0,
            sish=SimpleNamespace(
                shard_dir=str(sish_dir / "shards"),
                index_veb_path=str(sish_dir / "veb.pkl"),
                meta_database_path=str(sish_dir / "meta.pkl"),
            ),
        ),
        slide_retrieval=SimpleNamespace(exclusion_level="patient", search_workers=1),
    )
    task = SlideRetrievalTask(SimpleNamespace(cfg=cfg, project_root=str(tmp_path)))

    result = task.execute(
        combo_cfg=combo,
        datasets_by_use={"reference": [reference_dataset], "query": [query_dataset]},
    )

    assert result["num_queries"] == 1
    assert result["num_reference_items"] == 2
    assert (sish_dir / "veb.pkl").is_file()
    assert (sish_dir / "meta.pkl").is_file()

    from openpyxl import load_workbook

    workbook = load_workbook(Path(result["output_dir"]) / "query_results.xlsx", read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value) for value in rows[0]]
    first_result = dict(zip(headers, rows[1], strict=True))
    assert first_result["query_sample_id"] == "query-slide"
    assert first_result["rank_1_sample_id"] == "reference-match"
    assert float(first_result["rank_1_score"]) == pytest.approx(0.0)
